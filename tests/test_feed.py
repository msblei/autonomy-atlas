"""pipeline.feed on small cases (docs/methodology.md, "Feed")."""

import pandas as pd
import pytest

from pipeline import feed


def test_feed_table_is_plausible():
    items = feed.load_items()
    assert all(0 < v["protein"] < 0.7 for v in items.values())
    assert items[238]["parent"] == 236  # soy cake is made from soybeans
    assert all(v["parent"] is None or v["parent"] != k for k, v in items.items())


def test_eu_extra_shares_count_only_imports_from_outside():
    raw = pd.DataFrame([
        ("DEU", "BRA", 238, "imports", 2022, 300.0),
        ("DEU", "NLD", 238, "imports", 2022, 100.0),  # intra-EU
        ("FRA", "ARG", 238, "imports", 2022, 100.0),
        ("USA", "BRA", 238, "imports", 2022, 999.0),  # not a member
        ("DEU", "BRA", 238, "exports", 2022, 999.0),  # not an import
    ], columns=["reporter", "partner", "item_code", "flow", "year", "value"])
    s = feed.eu_extra_shares(raw).set_index(["item_code", "year"])["share"]
    assert s[(238, 2022)] == pytest.approx(0.8)


def _dgagri_sheet(path, sheet, h, i):
    """A marketing-year sheet laid out as DG AGRI's: title rows, header, sources; (J) names (H), (I)."""
    import openpyxl

    wb = openpyxl.load_workbook(path) if path.exists() else openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    ws = wb.create_sheet(sheet)
    ws.append(["Updated on 1 Jan 2026"])
    ws.append([sheet.replace("-", "/")])
    ws.append(["Protein source", "(A)", "(B)", "(C)", "(D)", "(E)", "(F)", "(G)",
               "EU total\n feed use\n(H) = (E) * (G)", "Feed use \nEU origin\n(I) = (F) * (G)",
               "% feed use of EU origin\n(J) = (I) / (H)"])
    for name, k in [("CEREALS (of which)", "CEREALS"), ("SOYA BEAN MEALS (of which)", "SOYA BEAN MEALS"),
                    ("RAPESEED MEALS (of which)", "RAPESEED MEALS"), ("SUNFLOWER MEALS (of which)", "SUNFLOWER MEALS"),
                    ("ROUGHAGE", "ROUGHAGE"), ("TOTAL", "TOTAL")]:
        ws.append([name, None, None, None, None, None, None, None, h[k], i[k], i[k] / h[k]])
    wb.save(path)


def test_dgagri_comparison_leaves_roughage_out_and_reads_columns_by_header(tmp_path):
    import pandas as pd

    from pipeline import feedcheck

    h = {"CEREALS": 15, "SOYA BEAN MEALS": 12, "RAPESEED MEALS": 4, "SUNFLOWER MEALS": 2, "ROUGHAGE": 30, "TOTAL": 70}
    i = {"CEREALS": 13, "SOYA BEAN MEALS": 0.5, "RAPESEED MEALS": 3, "SUNFLOWER MEALS": 1, "ROUGHAGE": 30, "TOTAL": 53}
    p = tmp_path / "fpb.xlsx"
    for y in ("2021-22", "2022-23"):
        _dgagri_sheet(p, y, h, i)
    dg = feedcheck.dgagri(p, (2021, 2022))
    assert dg.attrs["years"] == ["2021-22", "2022-23"]
    assert dg.at["CONCENTRATES", "protein"] == 40 and dg.at["CONCENTRATES", "imported"] == 17
    assert dg.at["SOYA BEAN MEALS", "imported"] == 11.5
    fao = pd.DataFrame({"item_code": [238, 15, 999], "protein_t": [12e6, 14e6, 10e6],
                        "imported_t": [11e6, 1e6, 4e6]})
    c = feedcheck.compare(dg, fao).set_index("source")
    assert c.at["All concentrate feed (without roughage)", "fao_imported_pct"] == 100 * 16 / 36
    assert c.at["Soya bean meal", "fao_protein_mt"] == 12
    assert c.at["Sunflower meal", "fao_protein_mt"] == 0


def test_marketing_years_start_in_the_window_years():
    from pipeline import feedcheck

    assert feedcheck.marketing_years((2021, 2023)) == ["2021-22", "2022-23", "2023-24"]
    assert feedcheck.marketing_years((2099, 2099)) == ["2099-00"]
